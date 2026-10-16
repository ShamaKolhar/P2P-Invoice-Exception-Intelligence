"""
Phase 6 - Excel exception log.

Builds outputs/excel/P2P_Exception_Log.xlsx from the pipeline's own parquet
outputs. Every number in the workbook traces back to prepare_data.py /
p2p_exceptions.py / ai_extract.py / ai_triage.py - nothing here is invented.

Sheets: Summary KPIs | Exception Log | AI Triage | AI Extraction Accuracy |
Vendor Hotspots | Rules | Data Quality
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import (
    CASE_FEATURES_PARQUET,
    EVENT_LOG_PARQUET,
    DATA_PROCESSED_DIR,
    OUTPUTS_DIR,
    OUTPUTS_EXCEL_DIR,
    COL_CASE_ID,
)

EXTRACTION_RESULTS_PATH = DATA_PROCESSED_DIR / "ai_extraction_results.parquet"
TRIAGE_RESULTS_PATH = DATA_PROCESSED_DIR / "ai_triage_results.parquet"

HEADER_FILL = PatternFill(start_color="2E5C8A", end_color="2E5C8A", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

FLAG_COLUMNS = [
    "E1_invoice_before_gr",
    "E2_price_change",
    "E3_quantity_change",
    "E4_payment_block",
    "E5_repeated_invoice_receipt",
    "E6_long_cycle_time",
]

RULES = [
    ("E1", "Invoice before goods receipt", "Invoice Receipt timestamp earlier than Goods Receipt timestamp, for cases in the '3-way match, invoice after GR' item category (a compliance breach)."),
    ("E2", "Price change after PO creation", "A 'Change Price' event exists anywhere in the case's event history."),
    ("E3", "Quantity change", "A 'Change Quantity' event exists anywhere in the case's event history."),
    ("E4", "Payment block set", "A 'Set Payment Block' or 'Remove Payment Block' event exists in the case's event history."),
    ("E5", "Invoice receipt cancelled or repeated", "More than one 'Record Invoice Receipt' event exists for the case."),
    ("E6", "Long cycle time", "Cycle time (Create PO Item to Clear Invoice) is above the 90th percentile for the case's own item category."),
]


def style_header_row(ws, n_cols):
    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize_columns(ws, df):
    for i, col in enumerate(df.columns, start=1):
        max_len = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).head(200)])
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 3, 45)


def write_dataframe_sheet(wb, sheet_name, df):
    ws = wb.create_sheet(sheet_name)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(["" if pd.isna(v) else v for v in row])
    if len(df) > 0:
        style_header_row(ws, len(df.columns))
        autosize_columns(ws, df)
    return ws


def write_key_value_sheet(wb, sheet_name, pairs):
    ws = wb.create_sheet(sheet_name)
    ws.append(["Metric", "Value"])
    for key, value in pairs:
        ws.append([key, value])
    style_header_row(ws, 2)
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 25
    return ws


def build_summary_kpis(case_features):
    total_cases = len(case_features)
    total_events = int(case_features["activity_count"].sum())
    exception_rate = round(100 * case_features["has_exception"].mean(), 2)
    total_net_worth = round(case_features["net_worth_eur"].sum(skipna=True), 2)
    exception_net_worth = round(
        case_features.loc[case_features["has_exception"], "net_worth_eur"].sum(skipna=True), 2
    )
    cycle = case_features["cycle_time_days"].dropna()

    pairs = [
        ("Total PO line items (cases)", total_cases),
        ("Total events", total_events),
        ("Overall exception rate (%)", exception_rate),
        ("Total net worth (EUR)", total_net_worth),
        ("Net worth in exception cases (EUR)", exception_net_worth),
        ("Mean cycle time (days)", round(cycle.mean(), 2) if len(cycle) else "n/a"),
        ("Median cycle time (days)", round(cycle.median(), 2) if len(cycle) else "n/a"),
        ("90th percentile cycle time (days)", round(cycle.quantile(0.90), 2) if len(cycle) else "n/a"),
    ]
    for flag in FLAG_COLUMNS:
        if flag in case_features.columns:
            pairs.append((f"{flag} - case count", int(case_features[flag].sum())))

    if EXTRACTION_RESULTS_PATH.exists():
        extraction_df = pd.read_parquet(EXTRACTION_RESULTS_PATH)
        norm_cols = [c for c in extraction_df.columns if c.endswith("_normalised_match")]
        avg_accuracy = round(extraction_df[norm_cols].mean().mean() * 100, 1)
        pairs.append(("AI extraction - documents processed", len(extraction_df)))
        pairs.append(("AI extraction - avg normalised accuracy (%)", avg_accuracy))
    else:
        pairs.append(("AI extraction", "not run (no API key / --skip-ai used)"))

    if TRIAGE_RESULTS_PATH.exists():
        triage_df = pd.read_parquet(TRIAGE_RESULTS_PATH)
        pairs.append(("AI triage - cases reviewed", len(triage_df)))
    else:
        pairs.append(("AI triage", "not run (no API key / --skip-ai used)"))

    return pairs


def build_rules_sheet(wb):
    ws = wb.create_sheet("Rules")
    ws.append(["Rule ID", "Name", "Logic"])
    skipped_path = OUTPUTS_DIR / "rules_skipped.txt"
    skipped_text = skipped_path.read_text(encoding="utf-8") if skipped_path.exists() else ""
    for rule_id, name, logic in RULES:
        ws.append([rule_id, name, logic])
    style_header_row(ws, 3)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 90
    note_row = ws.max_row + 2
    ws.cell(row=note_row, column=1, value="Resolution notes:")
    ws.cell(row=note_row + 1, column=1, value=skipped_text or "All rule inputs were found in the data.")
    return ws


def build_data_quality_sheet(wb):
    ws = wb.create_sheet("Data Quality")
    dq_path = OUTPUTS_DIR / "data_quality_summary.txt"
    text = dq_path.read_text(encoding="utf-8") if dq_path.exists() else "data_quality_summary.txt not found - run prepare_data.py first."
    ws.append(["Data Quality Summary"])
    ws["A1"].font = Font(bold=True, size=13)
    for line in text.split("\n"):
        ws.append([line])
    ws.column_dimensions["A"].width = 100
    return ws


def build_vendor_hotspots(case_features, vendor_col):
    if vendor_col is None or vendor_col not in case_features.columns:
        return pd.DataFrame(columns=["vendor", "case_count", "exception_count", "exception_rate_pct", "total_net_worth_eur", "exception_net_worth_eur"])
    grouped = case_features.groupby(vendor_col).agg(
        case_count=("has_exception", "size"),
        exception_count=("has_exception", "sum"),
        total_net_worth_eur=("net_worth_eur", "sum"),
    )
    grouped["exception_rate_pct"] = round(100 * grouped["exception_count"] / grouped["case_count"], 2)
    exception_value = case_features[case_features["has_exception"]].groupby(vendor_col)["net_worth_eur"].sum()
    grouped["exception_net_worth_eur"] = exception_value
    grouped = grouped.sort_values("exception_rate_pct", ascending=False).reset_index()
    grouped = grouped.rename(columns={vendor_col: "vendor"})
    return grouped[["vendor", "case_count", "exception_count", "exception_rate_pct", "total_net_worth_eur", "exception_net_worth_eur"]]


def build_extraction_accuracy_frame():
    if not EXTRACTION_RESULTS_PATH.exists():
        return None, None
    df = pd.read_parquet(EXTRACTION_RESULTS_PATH)
    fields = ["vendor", "address", "document_date", "total_amount"]
    summary_rows = []
    for field in fields:
        summary_rows.append({
            "field": field,
            "exact_match_rate_pct": round(df[f"{field}_exact_match"].mean() * 100, 1),
            "normalised_match_rate_pct": round(df[f"{field}_normalised_match"].mean() * 100, 1),
        })
    summary_df = pd.DataFrame(summary_rows)
    return summary_df, df


def main():
    OUTPUTS_EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    case_features = pd.read_parquet(CASE_FEATURES_PARQUET)

    vendor_col = next((c for c in case_features.columns if "vendor" in c.lower()), None)

    wb = Workbook()
    wb.remove(wb.active)

    write_key_value_sheet(wb, "Summary KPIs", build_summary_kpis(case_features))

    exception_log_cols = [c for c in case_features.columns if c != "activity_count"]
    write_dataframe_sheet(wb, "Exception Log", case_features[exception_log_cols])

    if TRIAGE_RESULTS_PATH.exists():
        triage_df = pd.read_parquet(TRIAGE_RESULTS_PATH)
        write_dataframe_sheet(wb, "AI Triage", triage_df)
    else:
        ws = wb.create_sheet("AI Triage")
        ws.append(["AI triage was not run for this build (no ANTHROPIC_API_KEY, or run_all.py was called with --skip-ai)."])

    summary_df, full_df = build_extraction_accuracy_frame()
    if summary_df is not None:
        ws = write_dataframe_sheet(wb, "AI Extraction Accuracy", summary_df)
        start_row = len(summary_df) + 4
        ws.cell(row=start_row - 1, column=1, value="Per-document results:").font = Font(bold=True)
        ws.append([])
        ws.append(list(full_df.columns))
        style_header_row_at = ws.max_row
        for row in full_df.itertuples(index=False):
            ws.append(["" if pd.isna(v) else v for v in row])
    else:
        ws = wb.create_sheet("AI Extraction Accuracy")
        ws.append(["AI extraction was not run for this build (no ANTHROPIC_API_KEY, or run_all.py was called with --skip-ai)."])

    vendor_hotspots = build_vendor_hotspots(case_features, vendor_col)
    write_dataframe_sheet(wb, "Vendor Hotspots", vendor_hotspots)

    build_rules_sheet(wb)
    build_data_quality_sheet(wb)

    out_path = OUTPUTS_EXCEL_DIR / "P2P_Exception_Log.xlsx"
    wb.save(out_path)
    print(f"Saved {out_path}")


if __name__ == "__main__":
    main()
